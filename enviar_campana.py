import os
import sys
from pathlib import Path

# Asegurar codificación UTF-8 en terminal de Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# Asegurar que el directorio del script esté en sys.path
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import time
import random
import smtplib
from datetime import datetime
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.image import MIMEImage
from email.utils import formataddr
import openpyxl

import config

# Constantes de columnas en la hoja 'Contactos' (Base 1-indexed)
HEADER_ROW = 4
DATA_START_ROW = 5

COL_ID = 1
COL_NEGOCIO = 2
COL_CATEGORIA = 3
COL_SECTOR = 4
COL_CORREO = 5
COL_TELEFONO = 6
COL_ASUNTO = 12
COL_CUERPO = 13
COL_ESTADO_CAMPANA = 16
COL_FECHA_ENVIO = 17
COL_RESULTADO = 18

def limpiar_correos(raw_email):
    """Separa y limpia múltiples correos separados por ';' o ','."""
    if not raw_email:
        return []
    texto = str(raw_email).replace(";", ",").replace("\n", ",")
    partes = [p.strip() for p in texto.split(",") if p.strip()]
    validos = [p for p in partes if "@" in p and "." in p]
    return validos

def probar_conexion_smtp(user, password):
    """Valida que las credenciales de Gmail funcionen correctamente."""
    try:
        server = smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=15)
        server.login(user, password)
        server.quit()
        return True, "Conexión exitosa a Gmail SMTP."
    except smtplib.SMTPAuthenticationError:
        return False, "Error de autenticación: Verifica tu correo y la Contraseña de Aplicación de 16 letras."
    except Exception as e:
        return False, f"Error al conectar con Gmail: {e}"

def enviar_correo(destinatarios, asunto, cuerpo_texto, user, password, sender_name):
    """Envía un correo electrónico profesional vía SMTP con HTML, texto plano y logo alojado en la nube."""
    if isinstance(destinatarios, str):
        destinatarios = [destinatarios]

    # Usamos multipart/alternative (CERO archivos adjuntos, CERO chips de logo.png en Gmail)
    msg = MIMEMultipart("alternative")
    msg["Subject"] = asunto
    msg["From"] = formataddr((sender_name, user))
    msg["To"] = ", ".join(destinatarios)

    # 1. Versión Texto Plano (Prioridad antispam y lectores simples)
    part_text = MIMEText(cuerpo_texto, "plain", "utf-8")
    msg.attach(part_text)

    # 2. Versión HTML profesional con el logo oficial alojado en tu web de NeXo
    # URL directa de tu web: https://ne-xo-page.vercel.app/logo.png (No genera chip de adjunto)
    LOGO_URL = "https://ne-xo-page.vercel.app/logo.png"

    # Formatear párrafos y destacar el saludo
    lineas = cuerpo_texto.strip().split("\n")
    if lineas:
        saludo = lineas[0]
        cuerpo_restante = "\n".join(lineas[1:]).strip()

        # Convertir dobles saltos en párrafos HTML limpios
        parrafos = [p.strip() for p in cuerpo_restante.split("\n\n") if p.strip()]
        html_parrafos = ""
        for p in parrafos:
            p_limpio = p.replace("\n", "<br>")
            html_parrafos += f'<p style="margin: 0 0 16px 0; font-size: 14.5px; line-height: 1.65; color: #334155;">{p_limpio}</p>'
    else:
        saludo = ""
        html_parrafos = f'<p style="margin: 0 0 16px 0; font-size: 14.5px; line-height: 1.65; color: #334155;">{cuerpo_texto.replace(chr(10), "<br>")}</p>'

    # Reemplazar enlaces para que sean interactivos y elegantes
    html_parrafos = html_parrafos.replace(
        "https://ne-xo-page.vercel.app",
        '<a href="https://ne-xo-page.vercel.app" style="color: #0284c7; font-weight: 600; text-decoration: none;">ne-xo-page.vercel.app</a>'
    )
    html_parrafos = html_parrafos.replace(
        "+57 322 206 7870",
        '<a href="https://wa.me/573222067870" style="color: #0284c7; font-weight: 600; text-decoration: none;">+57 322 206 7870</a>'
    )

    # Destacar palabras clave en la firma
    html_parrafos = html_parrafos.replace("Equipo NeXo", "<strong>Equipo NeXo</strong>")
    html_parrafos = html_parrafos.replace("Somos NeXo", "Somos <strong>NeXo</strong>")

    html_template = f'''<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
</head>
<body style="margin: 0; padding: 20px 10px; background-color: #f8fafc; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; -webkit-font-smoothing: antialiased;">
    <table align="center" border="0" cellpadding="0" cellspacing="0" width="100%" style="max-width: 580px; background-color: #ffffff; border-radius: 8px; border: 1px solid #e2e8f0; margin: 0 auto; overflow: hidden; box-shadow: 0 1px 3px rgba(0,0,0,0.05);">
        <!-- CABECERA CON LOGO NEXO -->
        <tr>
            <td align="center" style="padding: 28px 20px 20px 20px; background-color: #ffffff;">
                <a href="https://ne-xo-page.vercel.app" target="_blank" style="text-decoration: none;">
                    <img src="{LOGO_URL}" alt="NeXo" width="135" style="display: block; border: 0; outline: none; max-width: 135px; height: auto;" />
                </a>
            </td>
        </tr>
        <!-- CONTENIDO DEL MENSAJE -->
        <tr>
            <td style="padding: 10px 32px 30px 32px; background-color: #ffffff;">
                <h3 style="margin: 0 0 18px 0; font-size: 17px; font-weight: 700; color: #0f172a; line-height: 1.4;">
                    {saludo}
                </h3>
                {html_parrafos}
            </td>
        </tr>
    </table>
</body>
</html>'''

    part_html = MIMEText(html_template, "html", "utf-8")
    msg.attach(part_html)

    server = smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=20)
    server.login(user, password)
    server.sendmail(user, destinatarios, msg.as_string())
    server.quit()

def obtener_resumen_prospectos():
    """Analiza la hoja de Contactos y devuelve estadísticas."""
    if not config.EXCEL_PATH.exists():
        print(f"\n[!] Error: No se encontró el archivo {config.EXCEL_PATH}")
        return None

    wb = openpyxl.load_workbook(config.EXCEL_PATH, data_only=True)
    ws = wb["Contactos"]

    categorias = {}
    total_con_correo = 0
    total_enviados = 0
    total_pendientes = 0

    for r in range(DATA_START_ROW, ws.max_row + 1):
        correo = ws.cell(r, COL_CORREO).value
        cat = ws.cell(r, COL_CATEGORIA).value or "Sin Categoría"
        estado = ws.cell(r, COL_ESTADO_CAMPANA).value or ""

        correos_validos = limpiar_correos(correo)
        if correos_validos:
            total_con_correo += 1
            if cat not in categorias:
                categorias[cat] = {"total": 0, "pendientes": 0, "enviados": 0}
            categorias[cat]["total"] += 1

            if "enviado" in str(estado).lower():
                total_enviados += 1
                categorias[cat]["enviados"] += 1
            else:
                total_pendientes += 1
                categorias[cat]["pendientes"] += 1

    wb.close()
    return {
        "categorias": categorias,
        "total_con_correo": total_con_correo,
        "total_enviados": total_enviados,
        "total_pendientes": total_pendientes,
    }

def mostrar_resumen():
    """Muestra el reporte actual de prospectos en consola."""
    resumen = obtener_resumen_prospectos()
    if not resumen:
        return

    print("\n" + "=" * 65)
    print(" [i] RESUMEN DE PROSPECTOS - NeXo La Dorada")
    print("=" * 65)
    print(f" Total negocios con correo público: {resumen['total_con_correo']}")
    print(f" [OK] Correos ya enviados:          {resumen['total_enviados']}")
    print(f" [..] Correos pendientes de envío:  {resumen['total_pendientes']}")
    print("-" * 65)
    print(f"{'Categoría':<35} | {'Total':<6} | {'Pendientes':<10} | {'Enviados':<8}")
    print("-" * 65)
    for cat, datos in resumen["categorias"].items():
        print(f"{cat[:35]:<35} | {datos['total']:<6} | {datos['pendientes']:<10} | {datos['enviados']:<8}")
    print("=" * 65 + "\n")

def cargar_plantillas(wb):
    """Carga las plantillas dinámicas desde la hoja 'Plantillas'."""
    plantillas = {}
    if "Plantillas" not in wb.sheetnames:
        return plantillas
    ws = wb["Plantillas"]
    for r in range(6, ws.max_row + 1):
        cat = ws.cell(r, 1).value
        asunto = ws.cell(r, 2).value
        cuerpo = ws.cell(r, 3).value
        if cat and str(cat).strip():
            plantillas[str(cat).strip().lower()] = {
                "asunto": asunto,
                "cuerpo": cuerpo
            }
    return plantillas

def modo_prueba(password, categoria="Hoteles y hospedajes"):
    """Envía un correo de prueba a la propia cuenta del usuario usando la plantilla actual."""
    print("\n[*] --- MODO PRUEBA: ENVÍO A TU PROPIO CORREO ---")
    destinatario = config.GMAIL_USER
    print(f"Destino: {destinatario} (Tu cuenta de NeXo)")

    asunto = "Sistema de gestión hotelera a la medida para HOTEL DE PRUEBA | NeXo"
    cuerpo = """Hola, equipo de HOTEL DE PRUEBA:\n\nSomos NeXo, una empresa tecnológica de La Dorada..."""

    if config.EXCEL_PATH.exists():
        try:
            wb = openpyxl.load_workbook(config.EXCEL_PATH, data_only=True)
            plantillas = cargar_plantillas(wb)
            p = plantillas.get(str(categoria).strip().lower())
            if p and p.get("asunto") and p.get("cuerpo"):
                asunto = str(p["asunto"]).replace("{NEGOCIO}", "NEGOCIO DE PRUEBA")
                cuerpo = str(p["cuerpo"]).replace("{NEGOCIO}", "NEGOCIO DE PRUEBA")
            wb.close()
        except Exception:
            pass

    print("\nEnviando mensaje de prueba...")
    try:
        enviar_correo([destinatario], asunto, cuerpo, config.GMAIL_USER, password, config.SENDER_NAME)
        print("[OK] ¡Correo de prueba enviado con ÉXITO!")
        print("👉 Abre tu aplicación de Gmail o tu navegador y revisa cómo se ve el mensaje y la firma.")
    except Exception as e:
        print(f"[!] Error al enviar el correo de prueba: {e}")

def ejecutar_campana(categoria_filtro=None, limite_envios=None, password=None):
    """Ejecuta el envío de correos con control anti-spam y guardado en Excel."""
    wb = openpyxl.load_workbook(config.EXCEL_PATH)
    ws = wb["Contactos"]
    plantillas = cargar_plantillas(wb)

    # Identificar filas pendientes
    filas_a_enviar = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        correo_raw = ws.cell(r, COL_CORREO).value
        cat = ws.cell(r, COL_CATEGORIA).value
        estado = ws.cell(r, COL_ESTADO_CAMPANA).value or ""

        correos = limpiar_correos(correo_raw)
        if correos and "enviado" not in str(estado).lower():
            if categoria_filtro is None or (cat and cat.strip().lower() == categoria_filtro.strip().lower()):
                filas_a_enviar.append(r)

    total_pendientes = len(filas_a_enviar)
    if total_pendientes == 0:
        print("\n[OK] ¡No hay prospectos pendientes para enviar con ese criterio!")
        wb.close()
        return

    if limite_envios and limite_envios < total_pendientes:
        filas_a_enviar = filas_a_enviar[:limite_envios]

    print(f"\n[>] Se prepararon {len(filas_a_enviar)} prospectos para enviar en esta tanda.")
    print(f"[*] Pausa entre correos: {config.DELAY_MIN} a {config.DELAY_MAX} segundos (Protección Anti-Spam).")
    confirma = input("\n¿Deseas iniciar el envío real ahora? (s/n): ").strip().lower()
    if confirma != "s":
        print("[!] Envío cancelado por el usuario.")
        wb.close()
        return

    enviados_exito = 0
    errores = 0

    for i, r in enumerate(filas_a_enviar, 1):
        id_prospecto = ws.cell(r, COL_ID).value
        negocio = ws.cell(r, COL_NEGOCIO).value
        cat = ws.cell(r, COL_CATEGORIA).value
        correo_raw = ws.cell(r, COL_CORREO).value
        
        # Obtener asunto y cuerpo: si existe en la hoja Plantillas, se actualiza dinámicamente
        p = plantillas.get(str(cat).strip().lower()) if cat else None
        if p and p.get("asunto") and p.get("cuerpo"):
            asunto = str(p["asunto"]).replace("{NEGOCIO}", str(negocio).strip())
            cuerpo = str(p["cuerpo"]).replace("{NEGOCIO}", str(negocio).strip())
        else:
            asunto = ws.cell(r, COL_ASUNTO).value
            cuerpo = ws.cell(r, COL_CUERPO).value

        destinatarios = limpiar_correos(correo_raw)
        print(f"\n[{i}/{len(filas_a_enviar)}] Enviando a: {negocio} ({', '.join(destinatarios)})...")

        try:
            enviar_correo(destinatarios, asunto, cuerpo, config.GMAIL_USER, password, config.SENDER_NAME)
            ahora_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            # Actualizar Excel
            ws.cell(r, COL_ESTADO_CAMPANA).value = "Enviado"
            ws.cell(r, COL_FECHA_ENVIO).value = ahora_str
            ws.cell(r, COL_RESULTADO).value = f"Enviado con éxito a: {', '.join(destinatarios)}"
            wb.save(config.EXCEL_PATH)

            print(f"   [OK] Enviado y guardado en Excel ({ahora_str})")
            enviados_exito += 1

        except Exception as e:
            print(f"   [!] Error al enviar a {negocio}: {e}")
            ws.cell(r, COL_RESULTADO).value = f"Error: {e}"
            wb.save(config.EXCEL_PATH)
            errores += 1

        # Pausa anti-spam si aún quedan correos por enviar
        if i < len(filas_a_enviar):
            pausa = random.randint(config.DELAY_MIN, config.DELAY_MAX)
            print(f"   [...] Esperando {pausa} segundos antes del siguiente prospecto...", end="", flush=True)
            for seg in range(pausa, 0, -1):
                time.sleep(1)
                if seg % 15 == 0 or seg <= 5:
                    print(f" {seg}s...", end="", flush=True)
            print(" ¡Listo!")

    wb.close()
    print("\n" + "=" * 65)
    print(f"[OK] Tanda finalizada: {enviados_exito} enviados con éxito, {errores} errores.")
    print("=" * 65 + "\n")

def solicitar_password():
    """Obtiene la contraseña de aplicación desde .env o por teclado."""
    pwd = config.GMAIL_APP_PASSWORD
    if not pwd:
        print("\n[!] No se encontró la contraseña en el archivo .env.")
        pwd = input("👉 Por favor, ingresa tu Contraseña de Aplicación de 16 letras de Gmail: ").strip().replace(" ", "")
    return pwd

def menu_principal():
    print("""
=================================================================
               >> NeXo AUTOMATIZACIÓN DE PROSPECCIÓN <<          
                    La Dorada, Caldas · 2026                     
=================================================================
    """)
    mostrar_resumen()

    password = solicitar_password()
    print("\nComprobando conexión con Gmail...")
    ok, msg = probar_conexion_smtp(config.GMAIL_USER, password)
    if not ok:
        print(f"\n[!] {msg}")
        print("Recuerda generar la clave en: https://myaccount.google.com/apppasswords")
        return
    print(f"[OK] {msg}")

    while True:
        print("""
-----------------------------------------------------------------
 MENÚ DE OPCIONES:
 [1] Enviar CORREO DE PRUEBA a mi propia cuenta (Ver resultado)
 [2] Enviar tanda a HOTELES Y HOSPEDAJES (Recomendado)
 [3] Enviar tanda a SALUD Y ATENCIÓN MÉDICA
 [4] Enviar tanda a EDUCACIÓN E INSTITUCIONES
 [5] Enviar tanda a CUALQUIER OTRA CATEGORÍA / TODOS
 [6] Ver estadísticas actuales
 [7] Salir
-----------------------------------------------------------------
        """)
        opcion = input("Selecciona una opción (1-7): ").strip()

        if opcion == "1":
            modo_prueba(password)
        elif opcion == "2":
            limite = input("¿Cuántos hoteles deseas enviar hoy? (ej. 20 o presiona Enter para todos los pendientes): ").strip()
            limite_num = int(limite) if limite.isdigit() else None
            ejecutar_campana(categoria_filtro="Hoteles y hospedajes", limite_envios=limite_num, password=password)
        elif opcion == "3":
            ejecutar_campana(categoria_filtro="Salud y atención médica", password=password)
        elif opcion == "4":
            ejecutar_campana(categoria_filtro="Educación e instituciones", password=password)
        elif opcion == "5":
            cat = input("Escribe el nombre exacto de la categoría (o presiona Enter para TODOS los pendientes): ").strip()
            cat_filtro = cat if cat else None
            limite = input("Límite de correos en esta tanda (o Enter para todos): ").strip()
            limite_num = int(limite) if limite.isdigit() else None
            ejecutar_campana(categoria_filtro=cat_filtro, limite_envios=limite_num, password=password)
        elif opcion == "6":
            mostrar_resumen()
        elif opcion == "7":
            print("\n¡Hasta luego! Mucho éxito con NeXo.")
            break
        else:
            print("[!] Opción no válida. Intenta de nuevo.")

if __name__ == "__main__":
    menu_principal()
