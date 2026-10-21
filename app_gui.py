import os
import sys
import time
import random
import threading
from datetime import datetime
from pathlib import Path

# Configurar path para importar config y enviar_campana
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext
import openpyxl

import config
import enviar_campana

# Paleta de colores NeXo Dark Mode
BG_DARK = "#0f172a"        # Slate 900
BG_CARD = "#1e293b"        # Slate 800
BG_INPUT = "#334155"       # Slate 700
BORDER_COLOR = "#475569"   # Slate 600
TEXT_MAIN = "#f8fafc"      # Slate 50
TEXT_MUTED = "#94a3b8"     # Slate 400
ACCENT_BLUE = "#0284c7"    # Sky 600 (NeXo)
ACCENT_HOVER = "#0369a1"   # Sky 700
ACCENT_GREEN = "#10b981"   # Emerald 500
ACCENT_RED = "#ef4444"     # Red 500
ACCENT_YELLOW = "#f59e0b"  # Amber 500

class NexoAutomationApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("NeXo | Automatización de Prospección Comercial")
        self.geometry("860x720")
        self.minsize(800, 650)
        self.configure(bg=BG_DARK)

        # Variables de control
        self.is_running = False
        self.stop_requested = False
        self.worker_thread = None

        self.categoria_var = tk.StringVar(value="Hoteles y hospedajes")
        self.limite_var = tk.StringVar(value="20")
        self.delay_min_var = tk.StringVar(value=str(config.DELAY_MIN))
        self.delay_max_var = tk.StringVar(value=str(config.DELAY_MAX))

        self.setup_styles()
        self.create_widgets()
        self.actualizar_metricas()

        # Log inicial
        self.log("Sistema NeXo inicializado correctamente.", "INFO")
        self.log(f"Remitente configurado: {config.SENDER_NAME} <{config.GMAIL_USER}>", "INFO")
        self.log(f"Base de datos: {config.EXCEL_PATH.name}", "INFO")

    def setup_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")

        # Configuración general de ttk
        style.configure(".", background=BG_DARK, foreground=TEXT_MAIN, font=("Segoe UI", 9))
        style.configure("Card.TFrame", background=BG_CARD)

        # Estilo para Combobox
        style.configure(
            "TCombobox",
            fieldbackground=BG_INPUT,
            background=BG_CARD,
            foreground=TEXT_MAIN,
            arrowcolor=TEXT_MAIN,
            padding=5
        )
        style.map("TCombobox", fieldbackground=[("readonly", BG_INPUT)])

        # Estilo para Progressbar
        style.configure(
            "NeXo.Horizontal.TProgressbar",
            troughcolor=BG_CARD,
            background=ACCENT_BLUE,
            thickness=14
        )

    def create_widgets(self):
        # 1. HEADER
        header_frame = tk.Frame(self, bg=BG_DARK, padx=20, pady=15)
        header_frame.pack(fill=tk.X)

        title_lbl = tk.Label(
            header_frame,
            text="🚀 NeXo Software | Prospección Automatizada",
            font=("Segoe UI", 16, "bold"),
            bg=BG_DARK,
            fg=TEXT_MAIN
        )
        title_lbl.pack(anchor="w")

        subtitle_lbl = tk.Label(
            header_frame,
            text="Campaña comercial La Dorada, Caldas · Envío seguro de correos por categoría",
            font=("Segoe UI", 10),
            bg=BG_DARK,
            fg=TEXT_MUTED
        )
        subtitle_lbl.pack(anchor="w")

        # 2. CARDS DE MÉTRICAS (KPIs)
        kpi_frame = tk.Frame(self, bg=BG_DARK, padx=20)
        kpi_frame.pack(fill=tk.X, pady=(0, 15))

        self.card_total = self.crear_kpi_card(kpi_frame, "Total con Correo", "0", ACCENT_BLUE)
        self.card_total.pack(side=tk.LEFT, expand=True, fill=tk.BOTH, padx=(0, 10))

        self.card_pendientes = self.crear_kpi_card(kpi_frame, "Pendientes", "0", ACCENT_YELLOW)
        self.card_pendientes.pack(side=tk.LEFT, expand=True, fill=tk.BOTH, padx=5)

        self.card_enviados = self.crear_kpi_card(kpi_frame, "Enviados con Éxito", "0", ACCENT_GREEN)
        self.card_enviados.pack(side=tk.LEFT, expand=True, fill=tk.BOTH, padx=(10, 0))

        # 3. PANEL DE CONFIGURACIÓN Y CONTROLES
        ctrl_frame = tk.Frame(self, bg=BG_CARD, padx=15, pady=15, highlightbackground=BORDER_COLOR, highlightthickness=1)
        ctrl_frame.pack(fill=tk.X, padx=20, pady=(0, 15))

        # Fila 1: Filtros
        row1 = tk.Frame(ctrl_frame, bg=BG_CARD)
        row1.pack(fill=tk.X, pady=(0, 10))

        lbl_cat = tk.Label(row1, text="Categoría:", bg=BG_CARD, fg=TEXT_MAIN, font=("Segoe UI", 9, "bold"))
        lbl_cat.pack(side=tk.LEFT, padx=(0, 10))

        categorias_disponibles = [
            "Hoteles y hospedajes",
            "Salud y atención médica",
            "Educación e instituciones",
            "Belleza y cuidado personal",
            "Empresas y servicios",
            "Alimentos, tiendas y consumo",
            "TODAS LAS CATEGORÍAS"
        ]
        self.combo_cat = ttk.Combobox(
            row1,
            textvariable=self.categoria_var,
            values=categorias_disponibles,
            state="readonly",
            width=28,
            style="TCombobox"
        )
        self.combo_cat.pack(side=tk.LEFT, padx=(0, 20))

        lbl_lim = tk.Label(row1, text="Cantidad a enviar:", bg=BG_CARD, fg=TEXT_MAIN, font=("Segoe UI", 9, "bold"))
        lbl_lim.pack(side=tk.LEFT, padx=(0, 10))

        self.entry_lim = tk.Entry(
            row1,
            textvariable=self.limite_var,
            width=8,
            bg=BG_INPUT,
            fg=TEXT_MAIN,
            insertbackground=TEXT_MAIN,
            relief=tk.FLAT,
            font=("Segoe UI", 10),
            justify="center"
        )
        self.entry_lim.pack(side=tk.LEFT, padx=(0, 20))

        lbl_hint = tk.Label(row1, text="(ej. 20 para enviar en tandas)", bg=BG_CARD, fg=TEXT_MUTED, font=("Segoe UI", 8))
        lbl_hint.pack(side=tk.LEFT)

        # Fila 2: Botonera de Acción
        row2 = tk.Frame(ctrl_frame, bg=BG_CARD)
        row2.pack(fill=tk.X, pady=(5, 0))

        self.btn_iniciar = tk.Button(
            row2,
            text="🚀 Iniciar Envío",
            command=self.iniciar_campana_thread,
            bg=ACCENT_BLUE,
            fg=TEXT_MAIN,
            activebackground=ACCENT_HOVER,
            activeforeground=TEXT_MAIN,
            font=("Segoe UI", 10, "bold"),
            relief=tk.FLAT,
            padx=16,
            pady=6,
            cursor="hand2"
        )
        self.btn_iniciar.pack(side=tk.LEFT, padx=(0, 10))

        self.btn_detener = tk.Button(
            row2,
            text="⏹️ Detener",
            command=self.detener_campana,
            bg=ACCENT_RED,
            fg=TEXT_MAIN,
            activebackground="#dc2626",
            activeforeground=TEXT_MAIN,
            font=("Segoe UI", 10, "bold"),
            relief=tk.FLAT,
            padx=14,
            pady=6,
            state=tk.DISABLED,
            cursor="hand2"
        )
        self.btn_detener.pack(side=tk.LEFT, padx=(0, 10))

        self.btn_prueba = tk.Button(
            row2,
            text="🧪 Enviar Prueba a mi Correo",
            command=self.enviar_prueba_thread,
            bg=BG_INPUT,
            fg=TEXT_MAIN,
            activebackground=BORDER_COLOR,
            activeforeground=TEXT_MAIN,
            font=("Segoe UI", 9, "bold"),
            relief=tk.FLAT,
            padx=12,
            pady=6,
            cursor="hand2"
        )
        self.btn_prueba.pack(side=tk.LEFT, padx=(0, 10))

        self.btn_abrir_excel = tk.Button(
            row2,
            text="📂 Abrir Excel",
            command=self.abrir_archivo_excel,
            bg=BG_INPUT,
            fg=TEXT_MAIN,
            activebackground=BORDER_COLOR,
            activeforeground=TEXT_MAIN,
            font=("Segoe UI", 9),
            relief=tk.FLAT,
            padx=10,
            pady=6,
            cursor="hand2"
        )
        self.btn_abrir_excel.pack(side=tk.RIGHT)

        self.btn_recargar = tk.Button(
            row2,
            text="🔄 Recargar",
            command=self.actualizar_metricas,
            bg=BG_INPUT,
            fg=TEXT_MAIN,
            activebackground=BORDER_COLOR,
            activeforeground=TEXT_MAIN,
            font=("Segoe UI", 9),
            relief=tk.FLAT,
            padx=10,
            pady=6,
            cursor="hand2"
        )
        self.btn_recargar.pack(side=tk.RIGHT, padx=10)

        # 4. BARRA DE ESTADO Y PROGRESO
        status_frame = tk.Frame(self, bg=BG_DARK, padx=20)
        status_frame.pack(fill=tk.X, pady=(0, 8))

        self.lbl_status = tk.Label(
            status_frame,
            text="Listo para enviar.",
            bg=BG_DARK,
            fg=TEXT_MUTED,
            font=("Segoe UI", 9)
        )
        self.lbl_status.pack(anchor="w", pady=(0, 4))

        self.progress_bar = ttk.Progressbar(
            status_frame,
            style="NeXo.Horizontal.TProgressbar",
            mode="determinate"
        )
        self.progress_bar.pack(fill=tk.X)

        # 5. CONSOLA DE LOGS EN VIVO
        log_frame = tk.Frame(self, bg=BG_DARK, padx=20)
        log_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))

        lbl_log = tk.Label(log_frame, text="Registro de Actividad en Vivo:", bg=BG_DARK, fg=TEXT_MAIN, font=("Segoe UI", 9, "bold"))
        lbl_log.pack(anchor="w", pady=(0, 4))

        self.log_text = scrolledtext.ScrolledText(
            log_frame,
            bg=BG_CARD,
            fg=TEXT_MAIN,
            insertbackground=TEXT_MAIN,
            font=("Consolas", 9),
            relief=tk.FLAT,
            highlightbackground=BORDER_COLOR,
            highlightthickness=1,
            wrap=tk.WORD
        )
        self.log_text.pack(fill=tk.BOTH, expand=True)

        # Configurar colores de tags para el log
        self.log_text.tag_config("INFO", foreground="#38bdf8")     # Cyan
        self.log_text.tag_config("SUCCESS", foreground="#34d399")  # Green
        self.log_text.tag_config("WARN", foreground="#fbbf24")     # Yellow
        self.log_text.tag_config("ERROR", foreground="#f87171")    # Red
        self.log_text.tag_config("MUTED", foreground="#94a3b8")    # Gray

    def crear_kpi_card(self, parent, titulo, valor_inicial, color_acento):
        card = tk.Frame(parent, bg=BG_CARD, padx=15, pady=12, highlightbackground=BORDER_COLOR, highlightthickness=1)
        lbl_title = tk.Label(card, text=titulo.upper(), font=("Segoe UI", 8, "bold"), bg=BG_CARD, fg=TEXT_MUTED)
        lbl_title.pack(anchor="w")

        lbl_val = tk.Label(card, text=valor_inicial, font=("Segoe UI", 20, "bold"), bg=BG_CARD, fg=color_acento)
        lbl_val.pack(anchor="w")
        card.val_label = lbl_val
        return card

    def log(self, mensaje, nivel="INFO"):
        hora = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(tk.END, f"[{hora}] ", "MUTED")
        self.log_text.insert(tk.END, f"{mensaje}\n", nivel)
        self.log_text.see(tk.END)

    def actualizar_metricas(self):
        resumen = enviar_campana.obtener_resumen_prospectos()
        if resumen:
            self.card_total.val_label.config(text=str(resumen["total_con_correo"]))
            self.card_pendientes.val_label.config(text=str(resumen["total_pendientes"]))
            self.card_enviados.val_label.config(text=str(resumen["total_enviados"]))

    def abrir_archivo_excel(self):
        try:
            if config.EXCEL_PATH.exists():
                os.startfile(str(config.EXCEL_PATH))
                self.log(f"Abriendo archivo Excel: {config.EXCEL_PATH.name}", "INFO")
            else:
                messagebox.showerror("Error", f"No se encontró el archivo: {config.EXCEL_PATH}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo abrir el archivo: {e}")

    def enviar_prueba_thread(self):
        if self.is_running:
            return
        t = threading.Thread(target=self._enviar_prueba_worker, daemon=True)
        t.start()

    def _enviar_prueba_worker(self):
        self.btn_prueba.config(state=tk.DISABLED)
        self.lbl_status.config(text="Enviando correo de prueba...")
        self.log("Iniciando envío de prueba a tu propia cuenta...", "INFO")

        destinatario = config.GMAIL_USER
        cat_actual = self.categoria_var.get()
        cat_prueba = "Hoteles y hospedajes" if cat_actual == "TODAS LAS CATEGORÍAS" else cat_actual

        asunto = "Propuesta comercial para NEGOCIO DE PRUEBA | NeXo"
        cuerpo = "Hola, equipo de NEGOCIO DE PRUEBA:\n\nSomos NeXo, empresa tecnológica de La Dorada..."

        try:
            wb = openpyxl.load_workbook(config.EXCEL_PATH, data_only=True)
            plantillas = enviar_campana.cargar_plantillas(wb)
            p = plantillas.get(str(cat_prueba).strip().lower())
            if p and p.get("asunto") and p.get("cuerpo"):
                asunto = str(p["asunto"]).replace("{NEGOCIO}", "NEGOCIO DE PRUEBA")
                cuerpo = str(p["cuerpo"]).replace("{NEGOCIO}", "NEGOCIO DE PRUEBA")
            wb.close()
        except Exception:
            pass

        try:
            enviar_campana.enviar_correo(
                [destinatario],
                asunto,
                cuerpo,
                config.GMAIL_USER,
                config.GMAIL_APP_PASSWORD,
                config.SENDER_NAME
            )
            self.log(f"✅ Correo de prueba ({cat_prueba}) enviado con éxito a: {destinatario}", "SUCCESS")
            self.lbl_status.config(text="Correo de prueba enviado con éxito.")
            messagebox.showinfo("Éxito", f"Correo de prueba ({cat_prueba}) enviado a:\n{destinatario}\n\nRevisa tu bandeja de entrada.")
        except Exception as e:
            self.log(f"❌ Error al enviar prueba: {e}", "ERROR")
            self.lbl_status.config(text="Error al enviar prueba.")
            messagebox.showerror("Error", f"Fallo al conectar con Gmail:\n{e}")
        finally:
            self.btn_prueba.config(state=tk.NORMAL)

    def iniciar_campana_thread(self):
        if self.is_running:
            return

        cat_seleccionada = self.categoria_var.get()
        categoria_filtro = None if cat_seleccionada == "TODAS LAS CATEGORÍAS" else cat_seleccionada

        limite_str = self.limite_var.get().strip()
        limite_num = int(limite_str) if limite_str.isdigit() else None

        # Confirmación
        msg_confirm = f"¿Deseas iniciar el envío para:\n\nCategoría: {cat_seleccionada}\nLímite de tanda: {limite_num or 'Todos los pendientes'}?"
        if not messagebox.askyesno("Confirmar Envío", msg_confirm):
            return

        self.is_running = True
        self.stop_requested = False
        self.btn_iniciar.config(state=tk.DISABLED)
        self.btn_detener.config(state=tk.NORMAL)
        self.btn_prueba.config(state=tk.DISABLED)
        self.combo_cat.config(state=tk.DISABLED)

        self.worker_thread = threading.Thread(
            target=self._ejecutar_campana_worker,
            args=(categoria_filtro, limite_num),
            daemon=True
        )
        self.worker_thread.start()

    def detener_campana(self):
        if self.is_running:
            self.stop_requested = True
            self.lbl_status.config(text="Detención solicitada. Esperando terminar el correo actual...")
            self.log("⚠️ Solicitud de detención recibida. Se detendrá al finalizar el proceso actual.", "WARN")
            self.btn_detener.config(state=tk.DISABLED)

    def _ejecutar_campana_worker(self, categoria_filtro, limite_envios):
        try:
            wb = openpyxl.load_workbook(config.EXCEL_PATH)
            ws = wb["Contactos"]
            plantillas = enviar_campana.cargar_plantillas(wb)

            filas_a_enviar = []
            for r in range(enviar_campana.DATA_START_ROW, ws.max_row + 1):
                correo_raw = ws.cell(r, enviar_campana.COL_CORREO).value
                cat = ws.cell(r, enviar_campana.COL_CATEGORIA).value
                estado = ws.cell(r, enviar_campana.COL_ESTADO_CAMPANA).value or ""

                correos = enviar_campana.limpiar_correos(correo_raw)
                if correos and "enviado" not in str(estado).lower():
                    if categoria_filtro is None or (cat and cat.strip().lower() == categoria_filtro.strip().lower()):
                        filas_a_enviar.append(r)

            total_pendientes = len(filas_a_enviar)
            if total_pendientes == 0:
                self.log("No hay prospectos pendientes con el criterio seleccionado.", "WARN")
                self.lbl_status.config(text="Sin prospectos pendientes.")
                wb.close()
                return

            if limite_envios and limite_envios < total_pendientes:
                filas_a_enviar = filas_a_enviar[:limite_envios]

            total_tanda = len(filas_a_enviar)
            self.log(f"🚀 Iniciando tanda de {total_tanda} prospectos...", "INFO")
            self.progress_bar.config(maximum=total_tanda, value=0)

            enviados = 0
            errores = 0

            for i, r in enumerate(filas_a_enviar, 1):
                if self.stop_requested:
                    self.log("⏹️ Campaña detenida por el usuario.", "WARN")
                    break

                negocio = ws.cell(r, enviar_campana.COL_NEGOCIO).value or ""
                cat = ws.cell(r, enviar_campana.COL_CATEGORIA).value
                correo_raw = ws.cell(r, enviar_campana.COL_CORREO).value

                # Cargar plantilla dinámica si existe en la hoja Plantillas
                p = plantillas.get(str(cat).strip().lower()) if cat else None
                if p and p.get("asunto") and p.get("cuerpo"):
                    asunto = str(p["asunto"]).replace("{NEGOCIO}", str(negocio).strip())
                    cuerpo = str(p["cuerpo"]).replace("{NEGOCIO}", str(negocio).strip())
                else:
                    asunto = ws.cell(r, enviar_campana.COL_ASUNTO).value
                    cuerpo = ws.cell(r, enviar_campana.COL_CUERPO).value

                destinatarios = enviar_campana.limpiar_correos(correo_raw)

                self.lbl_status.config(text=f"Enviando [{i}/{total_tanda}]: {negocio}...")
                self.log(f"[{i}/{total_tanda}] Enviando a: {negocio} ({', '.join(destinatarios)})...", "INFO")

                try:
                    enviar_campana.enviar_correo(
                        destinatarios,
                        asunto,
                        cuerpo,
                        config.GMAIL_USER,
                        config.GMAIL_APP_PASSWORD,
                        config.SENDER_NAME
                    )
                    ahora_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                    ws.cell(r, enviar_campana.COL_ESTADO_CAMPANA).value = "Enviado"
                    ws.cell(r, enviar_campana.COL_FECHA_ENVIO).value = ahora_str
                    ws.cell(r, enviar_campana.COL_RESULTADO).value = f"Enviado con éxito a: {', '.join(destinatarios)}"
                    wb.save(config.EXCEL_PATH)

                    self.log(f"   ✅ Enviado y guardado en Excel ({ahora_str})", "SUCCESS")
                    enviados += 1
                except Exception as e:
                    self.log(f"   ❌ Error al enviar a {negocio}: {e}", "ERROR")
                    ws.cell(r, enviar_campana.COL_RESULTADO).value = f"Error: {e}"
                    wb.save(config.EXCEL_PATH)
                    errores += 1

                self.progress_bar.config(value=i)
                self.actualizar_metricas()

                # Pausa Anti-Spam
                if i < total_tanda and not self.stop_requested:
                    pausa = random.randint(config.DELAY_MIN, config.DELAY_MAX)
                    for seg in range(pausa, 0, -1):
                        if self.stop_requested:
                            break
                        self.lbl_status.config(text=f"Pausa anti-spam: esperando {seg}s antes del siguiente correo...")
                        time.sleep(1)

            wb.close()
            self.log(f"🎉 Tanda finalizada: {enviados} enviados con éxito, {errores} errores.", "SUCCESS")
            self.lbl_status.config(text=f"Finalizado: {enviados} enviados, {errores} errores.")
            messagebox.showinfo("Campaña Terminada", f"Se completó la tanda:\n\n• Enviados: {enviados}\n• Errores: {errores}")

        except Exception as e:
            self.log(f"Error general en la campaña: {e}", "ERROR")
            messagebox.showerror("Error", f"Ocurrió un error: {e}")
        finally:
            self.is_running = False
            self.stop_requested = False
            self.btn_iniciar.config(state=tk.NORMAL)
            self.btn_detener.config(state=tk.DISABLED)
            self.btn_prueba.config(state=tk.NORMAL)
            self.combo_cat.config(state="readonly")
            self.actualizar_metricas()

if __name__ == "__main__":
    app = NexoAutomationApp()
    app.mainloop()
